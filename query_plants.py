#!/usr/bin/env python3
"""
Plant habitat evidence scraper.
Output: plants_full_evidence.json

Sources (2 scrapers + 1 API, 2 tiers):
  Tier 1 — all genera (Burke + Go Botany run together, records from both kept):
    Burke Herbarium Image Collection 
      URL: https://burkeherbarium.org/imagecollection/taxon.php?Taxon={genus}
      Fields: Habitat field from species page (free text)
      Scrapes genus page → species links → each species page for "Habitat:" line.
      All species scraped (no max limit).
    Go Botany / Native Plant Trust 
      URL: https://gobotany.nativeplanttrust.org/species/{genus}/{species}/
      Fields: Habitat section (text), Characteristics Habitat (Type: terrestrial/
              aquatic/wetlands), Wetland indicator status (OBL/FACW/FAC/FACU)
      All species scraped (no max limit).
  Tier 2 — gap-filling (only genera with <3 unique species from Tier 1):
    NatureServe Explorer API (Canada only)
      URL: https://explorer.natureserve.org/api/data/speciesSearch
      Fields: majorHabitat (text), habitatComments (text)
      Filter: nation=CA, informalTaxonomy=Plants, kingdom=Plantae.
      Max 5 species per genus.

Pipeline:
  1. Read genera from Taxon list.xlsx (Group = "Plants")
  2. Tier 1: scrape Burke Herbarium + Go Botany for all species per genus
  3. Tier 2: genera with <3 unique species → query NatureServe Canada
  4. Save all results to JSON

Run log:
  Last run: 2026-09-07
  Environment: macOS 26.3, Python 3.13.13, openpyxl 3.1.5
  Result: 158/192 genera with evidence (Burke + Go Botany)
"""

import json
import time
import re
import urllib.request
import urllib.error
import urllib.parse
import openpyxl
from pathlib import Path

XLSX_FILE = Path(__file__).parent / "Taxon list.xlsx"
OUTPUT_JSON = Path(__file__).parent / "plants_full_evidence.json"
DELAY = 0.3


def extract_burke_all(genus):
    """Burke Herbarium: genus → all species with Habitat field."""
    url = f"https://burkeherbarium.org/imagecollection/taxon.php?Taxon={genus}"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        resp = urllib.request.urlopen(req, timeout=15)
        html = resp.read().decode("utf-8", errors="replace")
    except Exception:
        return []

    links = re.findall(r"Taxon=([^'\"&]+)", html)
    links = list(dict.fromkeys(links))
    species_list = [
        l for l in links
        if l.startswith(genus) and " " in l and "var." not in l and "ssp." not in l
    ]

    results = []
    for sp_name in species_list:
        sp_url = f"https://burkeherbarium.org/imagecollection/taxon.php?Taxon={sp_name.replace(' ', '+')}"
        try:
            req = urllib.request.Request(sp_url, headers={"User-Agent": "Mozilla/5.0"})
            resp = urllib.request.urlopen(req, timeout=15)
            sp_html = resp.read().decode("utf-8", errors="replace")
        except Exception:
            continue

        for line in sp_html.split("\n"):
            if "Habitat:" in line:
                clean = re.sub(r"<[^>]+>", "", line).strip()
                if clean and len(clean) > 10:
                    results.append({
                        "species": sp_name,
                        "habitat": clean,
                        "source": "Burke Herbarium",
                        "url": sp_url,
                    })
                    break
        time.sleep(0.2)

    return results


def extract_gobotany_all(genus):
    """Go Botany: genus → all species with Habitat + Characteristics Habitat + Wetland."""
    genus_url = f"https://gobotany.nativeplanttrust.org/genus/{genus.lower()}/"
    try:
        req = urllib.request.Request(genus_url, headers={"User-Agent": "Mozilla/5.0"})
        resp = urllib.request.urlopen(req, timeout=15)
        html = resp.read().decode("utf-8")
        species = re.findall(rf"/species/{genus.lower()}/([^/]+)/", html)
        species = list(dict.fromkeys(species))
    except Exception:
        return []

    if not species:
        return []

    results = []
    for sp in species:
        sp_url = f"https://gobotany.nativeplanttrust.org/species/{genus.lower()}/{sp}/"
        try:
            req = urllib.request.Request(sp_url, headers={"User-Agent": "Mozilla/5.0"})
            resp = urllib.request.urlopen(req, timeout=15)
            sp_html = resp.read().decode("utf-8")
        except Exception:
            continue

        parts = []

        hab_match = re.search(r'<div id="habitat"[^>]*>(.*?)</div>', sp_html, re.DOTALL)
        if hab_match:
            hab_text = re.sub(r"<[^>]+>", " ", hab_match.group(1))
            hab_text = re.sub(r"\s+", " ", hab_text).strip()
            hab_text = re.sub(r"^Habitat\s*", "", hab_text).strip()
            if hab_text:
                parts.append(f"Habitat: {hab_text[:200]}")

        chars_hab = re.search(r'<dt>Habitat</dt>\s*<dd>\s*(.*?)\s*</dd>', sp_html, re.DOTALL)
        if chars_hab:
            ch_text = re.sub(r"<[^>]+>", " ", chars_hab.group(1))
            ch_text = re.sub(r"\s+", " ", ch_text).strip()
            if ch_text:
                parts.append(f"Type: {ch_text[:150]}")

        wet_match = re.search(r'<div id="wetland"[^>]*>(.*?)</div>', sp_html, re.DOTALL)
        if wet_match:
            wet_text = re.sub(r"<[^>]+>", " ", wet_match.group(1))
            wet_text = re.sub(r"\s+", " ", wet_text).strip()
            wet_text = re.sub(r"^Wetland status\s*", "", wet_text).strip()
            if wet_text:
                parts.append(f"Wetland: {wet_text[:150]}")

        if parts:
            results.append({
                "species": f"{genus} {sp}",
                "habitat": " | ".join(parts),
                "source": "Go Botany",
                "url": sp_url,
            })
        time.sleep(0.2)

    return results


def extract_natureserve_canada(genus):
    """NatureServe Explorer: search genus in Canada, Plants only."""
    search_url = "https://explorer.natureserve.org/api/data/speciesSearch"
    payload = json.dumps({
        "criteriaType": "species",
        "textCriteria": [{"paramType": "quickSearch", "searchToken": genus}],
        "locationCriteria": [{"paramType": "nation", "nation": "CA"}],
        "speciesTaxonomyCriteria": [
            {"paramType": "informalTaxonomy", "informalTaxonomy": "Plants"}
        ],
    }).encode("utf-8")

    try:
        req = urllib.request.Request(
            search_url,
            data=payload,
            headers={"Content-Type": "application/json", "Accept": "application/json"},
        )
        resp = urllib.request.urlopen(req, timeout=20)
        data = json.loads(resp.read())
    except Exception:
        return []

    search_results = data.get("results", [])
    results = []

    for r in search_results:
        sci_name = r.get("scientificName", "")
        if not sci_name.startswith(genus):
            continue
        sg = r.get("speciesGlobal", {})
        if sg.get("kingdom") != "Plantae":
            continue
        uid = r.get("uniqueId")
        if not uid:
            continue

        detail_url = f"https://explorer.natureserve.org/api/data/taxon/{uid}"
        try:
            req = urllib.request.Request(
                detail_url, headers={"Accept": "application/json"}
            )
            resp = urllib.request.urlopen(req, timeout=15)
            detail = json.loads(resp.read())
        except Exception:
            continue

        eco = detail.get("ecologicalInformation", {})
        habitat = eco.get("majorHabitat")
        comments = eco.get("habitatComments", "")

        parts = []
        if habitat:
            parts.append(f"majorHabitat: {habitat}")
        if comments:
            parts.append(f"Comments: {comments[:200]}")

        if parts:
            page_url = (
                f"https://explorer.natureserve.org/Taxon/{uid}"
                f"/{sci_name.replace(' ', '_')}"
            )
            results.append({
                "species": sci_name,
                "habitat": " | ".join(parts),
                "source": "NatureServe Explorer",
                "url": page_url,
            })

        time.sleep(0.3)
        if len(results) >= 5:
            break

    return results


def main():
    wb = openpyxl.load_workbook(XLSX_FILE)
    ws = wb["Database"]

    genera = []
    for row in range(2, ws.max_row + 1):
        group = ws.cell(row=row, column=3).value
        if group and "plant" in str(group).lower():
            genus = str(ws.cell(row=row, column=9).value).strip()
            genera.append(genus)

    total = len(genera)
    print(f"Scraping {total} plant genera...", flush=True)
    print(f"Strategy: Burke + Go Botany (all spp, both kept) → NatureServe if <3 unique spp", flush=True)
    print(flush=True)

    stats = {"burke_only": 0, "gobotany_only": 0, "both": 0, "natureserve": 0, "manual": 0}
    full_evidence = {}

    for idx, genus in enumerate(genera):
        if (idx + 1) % 10 == 0 or idx == 0:
            print(
                f"  Progress: {idx+1}/{total} — "
                f"Burke:{stats['burke_only']} GoBotany:{stats['gobotany_only']} "
                f"Both:{stats['both']} NatureServe:{stats['natureserve']} "
                f"Manual:{stats['manual']}",
                flush=True,
            )

        burke_results = extract_burke_all(genus)
        time.sleep(DELAY)
        gobotany_results = extract_gobotany_all(genus)
        time.sleep(DELAY)

        has_burke = len(burke_results) > 0
        has_gobotany = len(gobotany_results) > 0

        all_records = burke_results + gobotany_results
        unique_spp = set(r["species"].lower() for r in all_records)
        n_unique = len(unique_spp)

        ns_results = []
        if n_unique < 3:
            ns_results = extract_natureserve_canada(genus)
            time.sleep(DELAY)
            all_records.extend(ns_results)
            unique_spp.update(r["species"].lower() for r in ns_results)
            n_unique = len(unique_spp)

        if all_records:
            full_evidence[genus] = {
                "unique_species": n_unique,
                "total_records": len(all_records),
                "burke_count": len(burke_results),
                "gobotany_count": len(gobotany_results),
                "natureserve_count": len(ns_results),
                "all_species": [
                    {"species": r["species"], "habitat": r["habitat"],
                     "source": r["source"], "url": r["url"]}
                    for r in all_records
                ],
            }

        if has_burke and has_gobotany:
            stats["both"] += 1
        elif has_burke:
            stats["burke_only"] += 1
        elif has_gobotany:
            stats["gobotany_only"] += 1
        elif ns_results:
            stats["natureserve"] += 1
        else:
            stats["manual"] += 1

        if all_records:
            short = all_records[0]["habitat"][:60]
            src = "Burke+GoBotany" if has_burke and has_gobotany else (
                "Burke" if has_burke else "GoBotany" if has_gobotany else "NatureServe")
            print(f"  [{src}] {genus} ({n_unique} unique spp, {len(all_records)} records): {short}...", flush=True)
        else:
            print(f"  [MANUAL] {genus}: no data in any source", flush=True)

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(full_evidence, f, ensure_ascii=False, indent=2)

    print(flush=True)
    print("=" * 60, flush=True)
    print(f"DONE. Output: {OUTPUT_JSON}", flush=True)
    print(f"Burke only: {stats['burke_only']}", flush=True)
    print(f"Go Botany only: {stats['gobotany_only']}", flush=True)
    print(f"Both Burke+GoBotany: {stats['both']}", flush=True)
    print(f"NatureServe needed: {stats['natureserve']}", flush=True)
    print(f"Manual lookup needed: {stats['manual']}", flush=True)
    print(f"Total: {total}", flush=True)


if __name__ == "__main__":
    main()
