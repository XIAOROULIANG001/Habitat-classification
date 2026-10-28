#!/usr/bin/env python3
"""
Decision tree — writes evidence from JSON files into
Taxon list.xlsx (Database sheet, columns K-O) for subsequent classification.


Decision logic:
  1. If ≥3 unique species in JSON → randomly sample 3 as representative evidence(random seed 42)
  2. If 1-2 unique species → use all available
  3. If no data → mark for manual lookup  
"""

import json
import random
import openpyxl
from openpyxl.styles import Font, PatternFill
from pathlib import Path

XLSX_FILE = Path(__file__).parent / "Taxon list.xlsx"

GROUPS = {
    "plant": {
        "json": Path(__file__).parent / "plants_full_evidence.json",
        "filter": lambda group: group.strip() == "Plants",
    },
    "alga": {
        "json": Path(__file__).parent / "algae_full_evidence.json",
        "filter": lambda group: group.strip() == "Algae",
    },
    "arthropod": {
        "json": Path(__file__).parent / "arthropods_full_evidence.json",
        "filter": lambda group: group.strip() == "Arthropods",
    },
    "fungi": {
        "json": Path(__file__).parent / "fungi_full_evidence.json",
        "filter": lambda group: group.strip() == "Fungi",
    },
    "vertebrate": {
        "json": Path(__file__).parent / "vertebrates_full_evidence.json",
        "filter": lambda group: group.strip() == "Vertebrates",
    },
    "protist": {
        "json": Path(__file__).parent / "protists_full_evidence.json",
        "filter": lambda group: group.strip() == "Protists",
    },
    "other_invertebrate": {
        "json": Path(__file__).parent / "other_invertebrates_full_evidence.json",
        "filter": lambda group: group.strip() == "Other_invertebrates",
    },
}

random.seed(42)


def deduplicate_by_species(records):
    """Deduplicate records by species name, keeping first occurrence."""
    seen = set()
    deduped = []
    for r in records:
        key = r["species"].lower()
        if key not in seen:
            seen.add(key)
            deduped.append(r)
    return deduped


def write_group(ws, group_name, group_config):
    """Write evidence for one taxonomic group into the worksheet."""
    json_file = group_config["json"]
    group_filter = group_config["filter"]

    if not json_file.exists():
        print(f"  Skipping {group_name}: {json_file.name} not found", flush=True)
        return 0, 0, 0

    with open(json_file, encoding="utf-8") as f:
        evidence = json.load(f)

    rows = []
    for row in range(2, ws.max_row + 1):
        group = ws.cell(row=row, column=3).value
        if group and group_filter(str(group)):
            genus = str(ws.cell(row=row, column=9).value).strip()
            rows.append((row, genus))

    written = 0
    manual = 0
    skipped = 0
    genus_cache = {}

    for row, genus in rows:
        existing_evidence = ws.cell(row=row, column=11).value
        if existing_evidence:
            skipped += 1
            continue

        if genus not in evidence:
            ws.cell(row=row, column=12, value="manual lookup needed")
            ws.cell(row=row, column=15, value="Xiaorou")
            manual += 1
            continue

        if genus not in genus_cache:
            data = evidence[genus]
            all_records = data["all_species"]
            unique_records = deduplicate_by_species(all_records)
            if len(unique_records) >= 3:
                selected = random.sample(unique_records, 3)
            else:
                selected = unique_records
            genus_cache[genus] = selected

        selected = genus_cache[genus]

        # K (col 11): Evidence
        evidence_parts = [f"[{s['species']}] {s['habitat']}" for s in selected]
        ws.cell(row=row, column=11, value="\n".join(evidence_parts))

        # M (col 13): Sources
        sources_parts = [f"{s['species']}: {s['source']}" for s in selected]
        ws.cell(row=row, column=13, value="\n".join(sources_parts))

        # N (col 14): Link
        urls = list(dict.fromkeys(s["url"] for s in selected))
        ws.cell(row=row, column=14, value="\n".join(urls))

        # O (col 15): Assigned by
        ws.cell(row=row, column=15, value="Xiaorou")

        written += 1

    return written, manual, skipped


def main():
    wb = openpyxl.load_workbook(XLSX_FILE)
    ws = wb["Database"]

    # Headers K-O
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    new_headers = {
        11: "Evidence",
        12: "Note",
        13: "Sources",
        14: "Link",
        15: "Assigned by",
    }
    for col, title in new_headers.items():
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
    for col in (16, 17):
        cell = ws.cell(row=1, column=col)
        cell.value = None
        cell.font = Font()
        cell.fill = PatternFill(fill_type=None)

    print("Writing evidence to Taxon list.xlsx...", flush=True)
    print(flush=True)

    for group_name, group_config in GROUPS.items():
        print(f"=== {group_name.upper()} ===", flush=True)
        written, manual, skipped = write_group(ws, group_name, group_config)
        print(f"  Evidence written: {written}", flush=True)
        print(f"  Manual lookup: {manual}", flush=True)
        if skipped:
            print(f"  Skipped (already filled): {skipped}", flush=True)
        print(flush=True)

    wb.save(XLSX_FILE)
    print(f"DONE. Saved to: {XLSX_FILE}", flush=True)


if __name__ == "__main__":
    main()
