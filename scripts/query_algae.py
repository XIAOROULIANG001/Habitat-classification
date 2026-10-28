#!/usr/bin/env python3
"""
Algae habitat evidence scraper.
Output: algae_full_evidence.json

Source:
  WoRMS REST API (single source, no gap-filling tier)
  URL: https://www.marinespecies.org/rest/AphiaRecordsByName/{genus}
  Fields: isMarine, isBrackish, isFreshwater, isTerrestrial (boolean flags)

Pipeline:
  1. Read algae genera from Taxon list.xlsx (Group = "Algae")
  2. For each genus, search WoRMS by name (filter kingdom=Plantae) → get AphiaID
  3. Query AphiaChildrenByAphiaID → collect species-level environment flags (max 10)
  4. If no species found but genus has environment flags → use genus-level fallback
  5. Save all results to JSON

Run log:
  Last run: 2026-09-08
  Environment: macOS 26.3, Python 3.13.13, openpyxl 3.1.5
  Result: 150/161 genera with evidence
"""

import json
import time
import urllib.request
import urllib.error
import openpyxl
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
XLSX_FILE = PROJECT_ROOT / "data" / "Taxon list.xlsx"
OUTPUT_JSON = PROJECT_ROOT / "data" / "evidence" / "algae_full_evidence.json"
DELAY = 0.5

WORMS_BASE = "https://www.marinespecies.org/rest"


def env_string(record):
    """Build environment string from WoRMS boolean flags."""
    parts = []
    if record.get("isMarine"):
        parts.append("marine")
    if record.get("isBrackish"):
        parts.append("brackish")
    if record.get("isFreshwater"):
        parts.append("freshwater")
    if record.get("isTerrestrial"):
        parts.append("terrestrial")
    return ", ".join(parts) if parts else ""


def worms_search_genus(genus):
    """Search WoRMS for a genus, return AphiaID and genus-level environment."""
    url = f"{WORMS_BASE}/AphiaRecordsByName/{genus}?like=false&marine_only=false"
    try:
        req = urllib.request.Request(url, headers={
            "Accept": "application/json", "User-Agent": "Mozilla/5.0"
        })
        resp = urllib.request.urlopen(req, timeout=15)
        data = json.loads(resp.read())
    except Exception:
        return None, ""

    if not data:
        return None, ""

    for r in data:
        if (r.get("rank") == "Genus" and r.get("status") == "accepted"
                and r.get("kingdom") == "Plantae"):
            return r["AphiaID"], env_string(r)
    for r in data:
        if r.get("kingdom") == "Plantae":
            return r.get("AphiaID"), env_string(r)
    return None, ""


def worms_get_species(aphia_id, genus, max_species=10):
    """Get up to max_species species under a genus AphiaID with environment data."""
    results = []
    offset = 1

    while True:
        url = (
            f"{WORMS_BASE}/AphiaChildrenByAphiaID/{aphia_id}"
            f"?marine_only=false&offset={offset}"
        )
        try:
            req = urllib.request.Request(url, headers={
                "Accept": "application/json", "User-Agent": "Mozilla/5.0"
            })
            resp = urllib.request.urlopen(req, timeout=15)
            children = json.loads(resp.read())
        except urllib.error.HTTPError as e:
            if e.code == 204:
                break
            break
        except Exception:
            break

        if not children:
            break

        for c in children:
            if c.get("rank") != "Species":
                continue
            env = env_string(c)
            if not env:
                continue
            sci_name = c.get("scientificname", "")
            if not sci_name.startswith(genus):
                continue

            worms_url = c.get("url", f"https://www.marinespecies.org/aphia.php?p=taxdetails&id={c['AphiaID']}")
            results.append({
                "species": sci_name,
                "habitat": f"Environment: {env}",
                "source": "WoRMS",
                "url": worms_url,
            })

            if len(results) >= max_species:
                return results

        if len(children) < 50:
            break
        offset += 50
        time.sleep(DELAY)

    return results


def main():
    wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
    ws = wb["Database"]

    genera = []
    for row in range(2, ws.max_row + 1):
        group = ws.cell(row=row, column=3).value
        if group and "alga" in str(group).lower():
            genus = str(ws.cell(row=row, column=9).value).strip()
            if genus not in genera:
                genera.append(genus)
    wb.close()

    total = len(genera)
    print(f"Scraping {total} algae genera from WoRMS...", flush=True)
    print(flush=True)

    stats = {"found": 0, "manual": 0}
    full_evidence = {}

    for idx, genus in enumerate(genera):
        if (idx + 1) % 10 == 0 or idx == 0:
            print(
                f"  Progress: {idx+1}/{total} — "
                f"Found:{stats['found']} Manual:{stats['manual']}",
                flush=True,
            )

        aphia_id, genus_env = worms_search_genus(genus)
        time.sleep(DELAY)

        if not aphia_id:
            stats["manual"] += 1
            print(f"  [MANUAL] {genus}: not found in WoRMS", flush=True)
            continue

        species_results = worms_get_species(aphia_id, genus)
        time.sleep(DELAY)

        if genus_env and not species_results:
            species_results = [{
                "species": genus,
                "habitat": f"Environment: {genus_env} (genus-level)",
                "source": "WoRMS",
                "url": f"https://www.marinespecies.org/aphia.php?p=taxdetails&id={aphia_id}",
            }]

        if species_results:
            unique_spp = set(r["species"].lower() for r in species_results)
            full_evidence[genus] = {
                "unique_species": len(unique_spp),
                "total_records": len(species_results),
                "worms_count": len(species_results),
                "genus_environment": genus_env,
                "all_species": [
                    {"species": r["species"], "habitat": r["habitat"],
                     "source": r["source"], "url": r["url"]}
                    for r in species_results
                ],
            }
            stats["found"] += 1
            short = species_results[0]["habitat"][:60]
            print(f"  [WoRMS] {genus} ({len(unique_spp)} spp): {short}...", flush=True)
        else:
            stats["manual"] += 1
            print(f"  [MANUAL] {genus}: found in WoRMS but no env data", flush=True)

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(full_evidence, f, ensure_ascii=False, indent=2)

    print(flush=True)
    print("=" * 60, flush=True)
    print(f"DONE. Output: {OUTPUT_JSON}", flush=True)
    print(f"Found: {stats['found']}/{total}", flush=True)
    print(f"Manual lookup needed: {stats['manual']}", flush=True)


if __name__ == "__main__":
    main()
